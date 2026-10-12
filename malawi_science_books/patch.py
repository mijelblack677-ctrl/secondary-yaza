# -*- coding: utf-8 -*-
"""
patch.py — Patch Method generator (v2)
Usage:
    python3 patch.py content/physics/form4/electricity/PHY_F4_ELEC_M1_Current.md
Produces:
    output/docx/PHY_F4_ELEC_M1_Current.docx
"""

import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------- COLORS ----------
BANNER_BLUE   = "1B3A6B"
BIG_IDEA      = "FFF9C4"
STRUGGLE      = "FCE4EC"
DEEP_EXPL     = "FFFFFF"
DIAGRAM_BG    = "F3E5F5"
KEY_POINTS    = "E8F5E9"
WORKED        = "E3F2FD"
CHECK         = "FFFDE7"
EXAM_Q        = "FFF8E1"
SOLUTIONS     = "F1F8E9"
MISTAKES      = "FFEBEE"
EXAM_TIP      = "FFF3E0"
REMEMBER_RED  = "C62828"
WHITE         = "FFFFFF"


def shade(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    pPr.append(shd)


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_box(doc, title, body_lines, fill_color, title_color=BANNER_BLUE, title_size=13):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill_color)

    if title:
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(title_size)
        r.font.color.rgb = RGBColor.from_string(title_color)
    else:
        cell.paragraphs[0].text = ""

    for line in body_lines:
        bp = cell.add_paragraph()
        bp.paragraph_format.space_after = Pt(4)
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = bp.add_run(part[2:-2]); r.bold = True
            else:
                bp.add_run(part)
    doc.add_paragraph()


def add_banner(doc, text, subtitle):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, BANNER_BLUE)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(WHITE)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.italic = True
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor.from_string("E0E0E0")
    doc.add_paragraph()


def add_remember_strip(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, REMEMBER_RED)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("REMEMBER  •  " + text)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(WHITE)
    doc.add_paragraph()


def add_heading(doc, text, color=BANNER_BLUE, size=15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string("0D47A1")
    return p


def add_md_paragraph(doc, text):
    """Handles **bold**, basic equation display, and bullet lines."""
    text = text.strip()
    if not text:
        return

    # Bullet
    if text.startswith("- "):
        p = doc.add_paragraph(style='List Bullet')
        text = text[2:]
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2]); r.bold = True
            else:
                p.add_run(part)
        return

    # Equation $$...$$
    if text.startswith('$$') and text.endswith('$$'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text[2:-2].strip())
        r.italic = True
        r.font.size = Pt(12)
        return

    p = doc.add_paragraph()
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)


def parse_frontmatter(text):
    meta = {}
    if text.startswith('---'):
        end = text.find('---', 3)
        block = text[3:end]
        for line in block.strip().splitlines():
            if ':' in line:
                k, v = line.split(':', 1)
                meta[k.strip()] = v.strip()
    return meta


def build_docx(md_path: Path, out_dir: Path):
    raw = md_path.read_text(encoding='utf-8')
    meta = parse_frontmatter(raw)

    body = raw
    if raw.startswith('---'):
        end = raw.find('---', 3)
        body = raw[end+3:]

    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    subtitle = f"{meta.get('series','')}  ·  {meta.get('subject','')}  ·  Form {meta.get('form','')}  ·  {meta.get('exam','')}  ·  Patch {meta.get('patch_id','')}"
    add_banner(doc, meta.get('title', 'Untitled'), subtitle)

    sections = re.split(r'\n# ', '\n' + body)
    for sec in sections:
        sec = sec.strip()
        if not sec:
            continue
        lines = sec.splitlines()
        title = lines[0].strip()
        content = '\n'.join(lines[1:]).strip()

        # Dispatch
        if title.startswith('The Big Idea'):
            add_heading(doc, '1. The Big Idea')
            add_box(doc, '', [p.strip() for p in content.split('\n\n') if p.strip()], BIG_IDEA)
            continue

        if title.startswith('Why Students Struggle'):
            add_heading(doc, '2. Why Students Struggle')
            add_box(doc, '', [p.strip() for p in content.split('\n\n') if p.strip()], STRUGGLE)
            continue

        if title.startswith('Deep Explanation'):
            add_heading(doc, '3. Deep Explanation')
            blocks = re.split(r'\n## ', '\n' + content)
            for blk in blocks:
                blk = blk.strip()
                if not blk:
                    continue
                blk_lines = blk.splitlines()
                sub = blk_lines[0].strip()
                subbody = '\n'.join(blk_lines[1:]).strip()
                if sub and not sub.startswith('#'):
                    add_subheading(doc, sub)
                for para in subbody.split('\n\n'):
                    para = para.strip()
                    if not para:
                        continue
                    for line in para.splitlines():
                        add_md_paragraph(doc, line)
                    doc.add_paragraph()
            continue

        if title.startswith('Diagrams'):
            add_heading(doc, '4. Diagrams')
            # Show diagram spec lines as a light purple box
            spec = [p.strip() for p in content.split('\n\n') if p.strip() and not p.strip().startswith('*(')]
            add_box(doc, 'Diagram Specifications (for illustrator)', spec, DIAGRAM_BG)
            continue

        if title.startswith('Key Points'):
            bullets = [l.lstrip('- ').strip() for l in content.splitlines() if l.strip().startswith('-')]
            add_box(doc, '5. Key Points', bullets, KEY_POINTS)
            continue

        if title.startswith('Worked Example'):
            add_heading(doc, '6. Worked Examples')
            blocks = re.split(r'\n## ', '\n' + content)
            for blk in blocks:
                blk = blk.strip()
                if not blk:
                    continue
                blk_lines = blk.splitlines()
                sub = blk_lines[0].strip()
                subbody = '\n'.join(blk_lines[1:]).strip()
                if sub and not sub.startswith('#'):
                    add_subheading(doc, sub)
                add_box(doc, '', [p.strip() for p in subbody.split('\n\n') if p.strip()], WORKED)
            continue

        if title.startswith('Check Yourself'):
            add_heading(doc, '7. Check Yourself')
            add_box(doc, '', [l.strip() for l in content.splitlines() if l.strip()], CHECK)
            continue

        if title.startswith('Exam-Style'):
            add_heading(doc, '8. Exam-Style Questions (MSCE)')
            add_box(doc, '', [p.strip() for p in content.split('\n\n') if p.strip()], EXAM_Q)
            continue

        if title.startswith('Full Worked Solutions'):
            add_heading(doc, '9. Full Worked Solutions')
            add_box(doc, '', [p.strip() for p in content.split('\n\n') if p.strip()], SOLUTIONS)
            continue

        if title.startswith('Common Mistakes'):
            add_heading(doc, '10. Common Mistakes')
            bullets = [l.lstrip('- ').strip() for l in content.splitlines() if l.strip().startswith('-')]
            add_box(doc, '', bullets, MISTAKES, title_color="B71C1C")
            continue

        if title.startswith('Exam Tip'):
            clean = [l.lstrip('> ').strip() for l in content.split('\n\n') if l.strip()]
            add_box(doc, '11. Exam Tip', clean, EXAM_TIP, title_color="E65100")
            continue

        if title.startswith('Remember'):
            add_remember_strip(doc, content.replace('**','').strip())
            continue

        # Fallback
        add_heading(doc, title)
        add_md_paragraph(doc, content)

    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / f"{meta.get('patch_id', md_path.stem)}.docx"
    doc.save(out_file)
    print(f"✅ Saved: {out_file}")
    return out_file


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python3 patch.py <path_to_md>")
        sys.exit(1)
    build_docx(Path(sys.argv[1]), Path("output/docx"))