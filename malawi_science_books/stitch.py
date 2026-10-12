"""
stitch.py — Merge all patches in a folder into one master .docx
Usage:
    python stitch.py content/physics/form4 output/docx/Physics_Form4_Complete.docx
"""

import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK


def stitch(folder: Path, out_file: Path):
    md_files = sorted(folder.glob("*.md"))
    if not md_files:
        print("No .md files found.")
        return

    from patch import build_docx
    master = Document()

    # Remove default empty paragraph
    for p in list(master.paragraphs):
        p._element.getparent().remove(p._element)

    for md in md_files:
        tmp_dir = Path("output/_tmp")
        docx_path = build_docx(md, tmp_dir)
        sub = Document(docx_path)
        # Copy body elements
        for element in sub.element.body:
            master.element.body.append(element)
        # Page break between topics
        p = master.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    out_file.parent.mkdir(parents=True, exist_ok=True)
    master.save(out_file)
    print(f"✅ Master book saved: {out_file}")


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python stitch.py <folder_of_md> <output_docx>")
        sys.exit(1)
    stitch(Path(sys.argv[1]), Path(sys.argv[2]))